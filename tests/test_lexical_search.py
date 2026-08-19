from __future__ import annotations

from pathlib import Path

from docx import Document

from document_finder.ingestion.pipeline import ingest_directory
from document_finder.search.lexical import search_lexical


def _create_doc(path: Path, title: str, sections: list[tuple[str, str]]) -> None:
    document = Document()
    document.core_properties.title = title
    for heading, text in sections:
        document.add_heading(heading, level=1)
        document.add_paragraph(text)
    document.save(path)


def _indexed_corpus(tmp_path: Path) -> Path:
    source = tmp_path / "source"
    source.mkdir()
    _create_doc(
        source / "Procurement Kit Process.docx",
        "Procurement Kit",
        [
            ("Procurement Kit Creation", "Create a procurement kit for a new release."),
            ("Procurement Kit Approval", "Approve the procurement kit before release."),
        ],
    )
    _create_doc(
        source / "Engineering Change.docx",
        "Change Management",
        [("Engineering Change Request", "Raise an engineering change request.")],
    )
    _create_doc(
        source / "Filename Match Guide.docx",
        "",
        [("General procedure", "This content deliberately has no matching document-name terms.")],
    )
    duplicate_a = source / "a"
    duplicate_b = source / "b"
    duplicate_a.mkdir()
    duplicate_b.mkdir()
    _create_doc(duplicate_a / "Same Name.docx", "", [("Release A", "shared release terminology")])
    _create_doc(duplicate_b / "Same Name.docx", "", [("Release B", "shared release terminology")])
    database = tmp_path / "index.sqlite3"
    summary = ingest_directory(source, database)
    assert summary.failed == 0
    return database


def test_exact_document_terminology_and_section_terminology(tmp_path: Path) -> None:
    database = _indexed_corpus(tmp_path)
    procurement = search_lexical("procurement kit", database_path=database)
    approval = search_lexical("approval", database_path=database)
    assert procurement[0]["filename"] == "Procurement Kit Process.docx"
    assert procurement[0]["document_id"]
    assert procurement[0]["chunk_id"]
    assert procurement[0]["page"] is None
    assert approval[0]["section"] == "Procurement Kit Approval"


def test_multiple_matching_chunks_return_one_document(tmp_path: Path) -> None:
    database = _indexed_corpus(tmp_path)
    results = search_lexical("procurement kit", database_path=database)
    assert [result["filename"] for result in results].count("Procurement Kit Process.docx") == 1


def test_duplicate_filename_prevention(tmp_path: Path) -> None:
    database = _indexed_corpus(tmp_path)
    results = search_lexical("shared release terminology", database_path=database)
    assert [result["filename"] for result in results] == ["Same Name.docx"]


def test_empty_query_returns_no_results(tmp_path: Path) -> None:
    assert search_lexical("   ", database_path=tmp_path / "missing.sqlite3") == []


def test_filename_is_a_low_weight_lexical_fallback(tmp_path: Path) -> None:
    database = _indexed_corpus(tmp_path)
    results = search_lexical("filename match", database_path=database)
    assert [result["filename"] for result in results] == ["Filename Match Guide.docx"]


def test_multiple_terms_require_all_terms(tmp_path: Path) -> None:
    database = _indexed_corpus(tmp_path)
    results = search_lexical("engineering change", database_path=database)
    assert [result["filename"] for result in results] == ["Engineering Change.docx"]
