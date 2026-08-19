from __future__ import annotations

import base64
import sqlite3
from pathlib import Path

from docx import Document

from document_finder.ingestion.chunking import build_chunks
from document_finder.ingestion.parsers import parse_docx
from document_finder.ingestion.pipeline import ingest_directory
from document_finder.ingestion.structure import extract_sections
from document_finder.storage.models import ParsedDocument, SourceBlock


def _save_document(path: Path, *, with_table: bool = False, title: str = "") -> None:
    doc = Document()
    doc.core_properties.title = title
    doc.add_heading("Main process", level=1)
    doc.add_paragraph("Create the process record.")
    doc.add_heading("Approval", level=2)
    doc.add_paragraph("Approve the record.", style="List Bullet")
    if with_table:
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Step"
        table.cell(0, 1).text = "Owner"
        table.cell(1, 0).text = "Create"
        table.cell(1, 1).text = "Engineering"
    doc.save(path)


def test_heading_hierarchy_and_section_extraction() -> None:
    parsed = ParsedDocument(
        filename="x.docx", source_path="x.docx", mime_type="docx", title=None, metadata={},
        native_text_characters=30, embedded_image_count=0, ocr_required=False, warnings=[],
        blocks=[
            SourceBlock(1, "paragraph", "Top", heading_level=1),
            SourceBlock(2, "paragraph", "top content"),
            SourceBlock(3, "paragraph", "Child", heading_level=2),
            SourceBlock(4, "paragraph", "child content"),
            SourceBlock(5, "paragraph", "Next", heading_level=1),
        ],
    )
    sections = extract_sections(parsed)
    assert [section.heading for section in sections] == ["Top", "Child", "Next"]
    assert sections[1].parent_local_id == sections[0].local_id
    assert sections[1].path == ["Top", "Child"]
    assert sections[0].blocks[0].text == "top content"
    assert sections[1].blocks[0].text == "child content"


def test_chunking_stays_within_section_and_splits_large_content() -> None:
    parsed = ParsedDocument(
        filename="x.docx", source_path="x.docx", mime_type="docx", title=None, metadata={},
        native_text_characters=500, embedded_image_count=0, ocr_required=False, warnings=[],
        blocks=[
            SourceBlock(1, "paragraph", "A", heading_level=1),
            SourceBlock(2, "paragraph", "one " * 40),
            SourceBlock(3, "paragraph", "two " * 40),
        ],
    )
    sections = extract_sections(parsed)
    chunks = build_chunks(sections, max_characters=120)
    assert len(chunks) >= 2
    assert all(chunk.section_local_id == sections[0].local_id for chunk in chunks)
    assert all(chunk.heading_path == ["A"] for chunk in chunks)
    assert all(chunk.page is None for chunk in chunks)


def test_docx_parser_extracts_tables_and_preserves_filename_and_metadata(tmp_path: Path) -> None:
    source = tmp_path / "Original Name.docx"
    _save_document(source, with_table=True, title="Process title")
    parsed = parse_docx(source, "Original Name.docx")
    assert parsed.filename == "Original Name.docx"
    assert parsed.metadata["title"] == "Process title"
    assert any(block.kind == "table" and "Step | Owner" in block.text for block in parsed.blocks)
    assert any(block.kind == "list" for block in parsed.blocks)


def test_empty_image_only_docx_is_marked_for_ocr(tmp_path: Path) -> None:
    image = tmp_path / "pixel.png"
    image.write_bytes(base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScL92wAAAABJRU5ErkJggg=="
    ))
    source = tmp_path / "Image only.docx"
    doc = Document()
    doc.add_picture(str(image))
    doc.save(source)
    parsed = parse_docx(source, "Image only.docx")
    assert parsed.native_text_characters == 0
    assert parsed.embedded_image_count == 1
    assert parsed.ocr_required is True
    assert parsed.warnings


def test_reingestion_skips_unchanged_and_reprocesses_changed_file(tmp_path: Path) -> None:
    source = tmp_path / "data"
    source.mkdir()
    document_path = source / "Workflow.docx"
    _save_document(document_path)
    database = tmp_path / "index.sqlite3"

    first = ingest_directory(source, database)
    second = ingest_directory(source, database)
    doc = Document(document_path)
    doc.add_paragraph("Changed procedure detail.")
    doc.save(document_path)
    third = ingest_directory(source, database)

    assert (first.processed, first.skipped, first.failed) == (1, 0, 0)
    assert (second.processed, second.skipped, second.failed) == (0, 1, 0)
    assert (third.processed, third.skipped, third.failed) == (1, 0, 0)
    connection = sqlite3.connect(database)
    try:
        assert connection.execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 1
        assert connection.execute("SELECT filename FROM documents").fetchone()[0] == "Workflow.docx"
        assert connection.execute("SELECT COUNT(*) FROM chunks").fetchone()[0] > 0
        assert connection.execute("SELECT COUNT(*) FROM chunks_fts").fetchone()[0] > 0
    finally:
        connection.close()
