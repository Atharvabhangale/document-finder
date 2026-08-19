from __future__ import annotations

from pathlib import Path

from docx import Document

from document_finder.ingestion.pipeline import ingest_directory
from document_finder.search.hybrid import search_hybrid


def _doc(path: Path, heading: str, text: str) -> None:
    document = Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(text)
    document.save(path)


def test_hybrid_combines_vector_lexical_heading_and_deduplicates(tmp_path: Path) -> None:
    source = tmp_path / "source"
    (source / "a").mkdir(parents=True)
    (source / "b").mkdir()
    _doc(source / "Target.docx", "Adding BOM Structure", "procedure text")
    _doc(source / "a" / "Same.docx", "Change one", "change text")
    _doc(source / "b" / "Same.docx", "Change two", "change text")
    database = tmp_path / "index.sqlite3"
    ingest_directory(source, database)

    def lexical(query: str, limit: int, db: Path) -> list[dict[str, object]]:
        return [{"document_id": "lex", "filename": "Lexical.docx", "score": 1.0, "section": "L", "page": None, "chunk_id": "1"}]

    def vector(query: str, limit: int, db: Path) -> list[dict[str, object]]:
        return [{"document_id": "vec", "filename": "Vector.docx", "score": 1.0, "section": "V", "page": None, "chunk_id": "2"}]

    results = search_hybrid("add BOM structure", database_path=database, lexical_search=lexical, vector_search=vector)
    assert results[0]["filename"] == "Target.docx"  # generic heading-token evidence
    assert {item["filename"] for item in results}.issuperset({"Lexical.docx", "Vector.docx"})
    assert search_hybrid("", database_path=database) == []


def test_hybrid_keeps_vector_only_lexical_only_and_multiple_documents(tmp_path: Path) -> None:
    source = tmp_path / "source"
    source.mkdir()
    _doc(source / "Neutral.docx", "Neutral heading", "neutral text")
    database = tmp_path / "index.sqlite3"
    ingest_directory(source, database)

    def lexical(query: str, limit: int, db: Path) -> list[dict[str, object]]:
        return [
            {"document_id": "lexical-only", "filename": "Lexical.docx", "score": 1.0, "section": "L", "page": None, "chunk_id": "1"},
            {"document_id": "shared", "filename": "Shared.docx", "score": 0.9, "section": "L2", "page": None, "chunk_id": "2"},
        ]

    def vector(query: str, limit: int, db: Path) -> list[dict[str, object]]:
        return [
            {"document_id": "vector-only", "filename": "Vector.docx", "score": 1.0, "section": "V", "page": None, "chunk_id": "3"},
            {"document_id": "shared", "filename": "Shared.docx", "score": 0.9, "section": "V2", "page": None, "chunk_id": "4"},
        ]

    results = search_hybrid("unmatched", database_path=database, lexical_search=lexical, vector_search=vector)
    filenames = [result["filename"] for result in results]
    assert set(filenames) == {"Lexical.docx", "Vector.docx", "Shared.docx"}
    assert filenames.count("Shared.docx") == 1
