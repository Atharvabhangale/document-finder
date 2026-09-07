"""Corpus root configuration, document identity, and safe opening (Phase 15).

Temporary folders and a stub embedding model throughout: these never read the
36-document corpus and never download a model.
"""

from __future__ import annotations

import sqlite3
from pathlib import Path

import numpy as np
import pytest
from docx import Document
from fastapi.testclient import TestClient

from document_finder import config
from document_finder.api import routes
from document_finder.api.app import app
from document_finder.corpus import DOCUMENT_ROOT_KEY, refresh_corpus
from document_finder.search.vector import search_vector


client = TestClient(app)


class StubEmbeddingModel:
    model_name = "test/stub-embedding"
    configuration = {"version": "1", "normalize": True}

    @staticmethod
    def _vector(text: str) -> list[float]:
        text = text.casefold()
        return [float("alpha" in text), float("beta" in text), 0.1]

    def embed_documents(self, texts: list[str]) -> np.ndarray:
        return np.asarray([self._vector(text) for text in texts], dtype=np.float32)

    def embed_queries(self, texts: list[str]) -> np.ndarray:
        return np.asarray([self._vector(text) for text in texts], dtype=np.float32)


def make_docx(path: Path, heading: str, body: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(body)
    document.save(path)


def serve(monkeypatch: pytest.MonkeyPatch, root: Path) -> StubEmbeddingModel:
    """Point the application at `root` and index it, exactly as a demo would."""
    monkeypatch.setenv(config.DATA_ROOT_VARIABLE, str(root))
    monkeypatch.delenv(config.DATABASE_VARIABLE, raising=False)
    monkeypatch.delenv(config.INDEX_VARIABLE, raising=False)
    model = StubEmbeddingModel()
    refresh_corpus(root, model=model)
    # The API's default embedding model would need to be downloaded; route search
    # through the same production code path with the stub instead.
    monkeypatch.setattr(
        routes, "search_vector",
        lambda query, **kwargs: search_vector(
            query, kwargs.get("limit", 5), config.database_path(), config.index_path(), model
        ),
    )
    return model


# --------------------------------------------------------------------- configuration


def test_default_configuration_is_unchanged(monkeypatch: pytest.MonkeyPatch) -> None:
    """With nothing set, the previous development layout still applies."""
    for variable in (config.DATA_ROOT_VARIABLE, config.DATABASE_VARIABLE, config.INDEX_VARIABLE):
        monkeypatch.delenv(variable, raising=False)
    assert config.data_root() == Path("data")
    assert config.database_path() == Path("data/document_finder.sqlite3")
    assert config.index_path() == Path("data/vector/qwen3_embedding_0_6b.faiss")


def test_custom_corpus_root_moves_every_artifact(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    monkeypatch.setenv(config.DATA_ROOT_VARIABLE, str(tmp_path / "customer"))
    monkeypatch.delenv(config.DATABASE_VARIABLE, raising=False)
    monkeypatch.delenv(config.INDEX_VARIABLE, raising=False)
    assert config.data_root() == tmp_path / "customer"
    assert config.database_path() == tmp_path / "customer" / "document_finder.sqlite3"
    assert config.index_path().parent == tmp_path / "customer" / "vector"


def test_database_and_index_can_be_overridden_independently(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    monkeypatch.setenv(config.DATA_ROOT_VARIABLE, str(tmp_path / "docs"))
    monkeypatch.setenv(config.DATABASE_VARIABLE, str(tmp_path / "state" / "index.sqlite3"))
    monkeypatch.setenv(config.INDEX_VARIABLE, str(tmp_path / "state" / "vectors.faiss"))
    assert config.database_path() == tmp_path / "state" / "index.sqlite3"
    assert config.index_path() == tmp_path / "state" / "vectors.faiss"


def test_nonexistent_corpus_root_fails_clearly(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    monkeypatch.setenv(config.DATA_ROOT_VARIABLE, str(tmp_path / "not-there"))
    monkeypatch.delenv(config.DATABASE_VARIABLE, raising=False)

    assert routes.corpus_state() == ("unindexed", 0)
    response = client.post("/search", json={"query": "alpha"})
    assert response.status_code == 503
    assert "python -m document_finder.corpus" in response.json()["detail"]
    assert client.get("/health").json()["corpus"] == "unindexed"


def test_index_built_for_another_folder_is_refused(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """Indexing folder A and serving folder B must never silently succeed."""
    folder_a, folder_b = tmp_path / "a", tmp_path / "b"
    make_docx(folder_a / "Alpha.docx", "Alpha", "alpha content")
    make_docx(folder_b / "Beta.docx", "Beta", "beta content")
    refresh_corpus(folder_a, model=StubEmbeddingModel())

    # Serve folder B while pointing at folder A's index.
    monkeypatch.setenv(config.DATA_ROOT_VARIABLE, str(folder_b))
    monkeypatch.setenv(config.DATABASE_VARIABLE, str(folder_a / "document_finder.sqlite3"))

    state, _ = routes.corpus_state()
    assert state == "mismatch"
    response = client.post("/search", json={"query": "alpha"})
    assert response.status_code == 503
    assert "different document folder" in response.json()["detail"]
    assert client.get("/health").json()["corpus"] == "mismatch"


def test_indexing_records_the_folder_it_described(tmp_path: Path) -> None:
    source = tmp_path / "corpus"
    make_docx(source / "Alpha.docx", "Alpha", "alpha content")
    refresh_corpus(source, model=StubEmbeddingModel())

    connection = sqlite3.connect(source / "document_finder.sqlite3")
    try:
        stored = connection.execute(
            "SELECT value FROM index_metadata WHERE key = ?", (DOCUMENT_ROOT_KEY,)
        ).fetchone()[0]
    finally:
        connection.close()
    assert Path(stored) == source.resolve()


def test_legacy_index_without_recorded_root_still_serves(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """An index written before this metadata existed must keep working."""
    source = tmp_path / "corpus"
    make_docx(source / "Alpha.docx", "Alpha", "alpha content")
    refresh_corpus(source, model=StubEmbeddingModel())
    database = source / "document_finder.sqlite3"
    connection = sqlite3.connect(database)
    with connection:
        connection.execute("DELETE FROM index_metadata")
    connection.close()

    monkeypatch.setenv(config.DATA_ROOT_VARIABLE, str(source))
    monkeypatch.delenv(config.DATABASE_VARIABLE, raising=False)
    assert routes.corpus_state() == ("ok", 1)


# ------------------------------------------------------------------------- identity


def test_duplicate_filenames_stay_distinct_through_the_api(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    source = tmp_path / "corpus"
    make_docx(source / "Engineering" / "procedure.docx", "Alpha engineering", "alpha engineering procedure")
    make_docx(source / "Quality" / "procedure.docx", "Alpha quality", "alpha quality procedure")
    serve(monkeypatch, source)

    results = client.post("/search", json={"query": "alpha", "limit": 5}).json()["results"]

    assert len(results) == 2
    assert {result["filename"] for result in results} == {"procedure.docx"}
    assert {result["folder"] for result in results} == {"Engineering", "Quality"}
    assert len({result["document_id"] for result in results}) == 2
    assert all(len(result["document_id"]) == 64 for result in results)
    # No absolute machine path is ever exposed.
    assert all(not result["folder"].startswith("/") for result in results)
    assert all(str(tmp_path) not in str(result) for result in results)


def test_each_result_identifier_opens_its_own_document(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    source = tmp_path / "corpus"
    make_docx(source / "Engineering" / "procedure.docx", "Alpha engineering", "alpha engineering procedure")
    make_docx(source / "Quality" / "procedure.docx", "Alpha quality", "alpha quality procedure")
    serve(monkeypatch, source)
    results = client.post("/search", json={"query": "alpha", "limit": 5}).json()["results"]

    for result in results:
        opened = client.get(f"/documents/by-id/{result['document_id']}")
        assert opened.status_code == 200
        expected = (source / result["folder"] / result["filename"]).read_bytes()
        assert opened.content == expected, "identifier must open the exact document, not its namesake"


# -------------------------------------------------------------------------- opening


def test_valid_identifier_opens_from_the_configured_root(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    source = tmp_path / "corpus"
    make_docx(source / "Alpha Process.docx", "Alpha", "alpha content")
    serve(monkeypatch, source)
    results = client.post("/search", json={"query": "alpha", "limit": 5}).json()["results"]

    identifier = results[0]["document_id"]
    resolved = routes.resolve_document_id(identifier)
    assert resolved == (source / "Alpha Process.docx").resolve()
    assert client.get(f"/documents/by-id/{identifier}").status_code == 200


def test_unknown_and_malformed_identifiers_are_rejected(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    source = tmp_path / "corpus"
    make_docx(source / "Alpha.docx", "Alpha", "alpha content")
    serve(monkeypatch, source)

    assert routes.resolve_document_id("0" * 64) is None
    assert client.get(f"/documents/by-id/{'0' * 64}").status_code == 404
    for malformed in ("", "not-hex", "../../etc/passwd", "a" * 63, "A" * 64):
        assert routes.resolve_document_id(malformed) is None


def test_missing_source_file_is_handled_safely(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """A document deleted from disk since indexing must 404, not raise."""
    source = tmp_path / "corpus"
    make_docx(source / "Alpha.docx", "Alpha", "alpha content")
    serve(monkeypatch, source)
    identifier = client.post("/search", json={"query": "alpha"}).json()["results"][0]["document_id"]

    (source / "Alpha.docx").unlink()

    assert routes.resolve_document_id(identifier) is None
    assert client.get(f"/documents/by-id/{identifier}").status_code == 404


def test_traversal_via_a_stored_path_cannot_escape_the_root(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """Even a poisoned stored source path must not resolve outside the corpus."""
    source = tmp_path / "corpus"
    outside = tmp_path / "outside.docx"
    outside.write_bytes(b"must not be served")
    make_docx(source / "Alpha.docx", "Alpha", "alpha content")
    serve(monkeypatch, source)

    database = config.database_path()
    connection = sqlite3.connect(database)
    with connection:
        connection.execute(
            "UPDATE documents SET source_path = ? WHERE filename = ?", ("../outside.docx", "Alpha.docx")
        )
    identifier = connection.execute("SELECT document_id FROM documents").fetchone()[0]
    connection.close()

    assert routes.resolve_document_id(identifier) is None
    assert client.get(f"/documents/by-id/{identifier}").status_code == 404
    assert routes.resolve_document_path("Alpha.docx") is None


def test_absolute_stored_path_is_also_refused(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    source = tmp_path / "corpus"
    outside = tmp_path / "outside.docx"
    outside.write_bytes(b"must not be served")
    make_docx(source / "Alpha.docx", "Alpha", "alpha content")
    serve(monkeypatch, source)

    connection = sqlite3.connect(config.database_path())
    with connection:
        connection.execute("UPDATE documents SET source_path = ?", (str(outside),))
    identifier = connection.execute("SELECT document_id FROM documents").fetchone()[0]
    connection.close()

    assert routes.resolve_document_id(identifier) is None


# --------------------------------------------------------------------- compatibility


def test_single_filename_workflow_still_works(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """The pre-existing filename route keeps working for unambiguous names."""
    source = tmp_path / "corpus"
    make_docx(source / "Alpha Process.docx", "Alpha", "alpha content")
    serve(monkeypatch, source)

    assert routes.resolve_document_path("Alpha Process.docx") == (source / "Alpha Process.docx").resolve()
    assert client.get("/documents/Alpha%20Process.docx").status_code == 200
    assert client.get("/documents/Missing.docx").status_code == 404
