"""Corpus loading and refresh behaviour (Phase 14).

These use temporary folders and a stub embedding model, so they never touch the
36-document corpus and never download a model.
"""

from __future__ import annotations

import sqlite3
from pathlib import Path

import numpy as np
import pytest
from docx import Document

from document_finder.corpus import _prunable, index_consistency, refresh_corpus
from document_finder.search.vector import search_vector
from document_finder.storage.repository import SQLiteRepository


class StubEmbeddingModel:
    """Deterministic three-dimensional stand-in; counts how much work it is asked to do."""

    model_name = "test/stub-embedding"
    configuration = {"version": "1", "normalize": True}

    def __init__(self) -> None:
        self.embedded_texts = 0

    @staticmethod
    def _vector(text: str) -> list[float]:
        text = text.casefold()
        return [float("alpha" in text), float("beta" in text), 0.1]

    def embed_documents(self, texts: list[str]) -> np.ndarray:
        self.embedded_texts += len(texts)
        return np.asarray([self._vector(text) for text in texts], dtype=np.float32)

    def embed_queries(self, texts: list[str]) -> np.ndarray:
        return np.asarray([self._vector(text) for text in texts], dtype=np.float32)


def make_docx(path: Path, heading: str, body: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(body)
    document.save(path)


def make_pdf(path: Path, text: str) -> None:
    import fitz

    path.parent.mkdir(parents=True, exist_ok=True)
    document = fitz.open()
    document.new_page().insert_text((72, 72), text)
    document.save(path)
    document.close()


@pytest.fixture
def corpus(tmp_path: Path) -> tuple[Path, Path, Path, StubEmbeddingModel]:
    source = tmp_path / "corpus"
    make_docx(source / "Alpha Process.docx", "Alpha heading", "alpha process content")
    make_docx(source / "nested" / "Beta Process.docx", "Beta heading", "beta process content")
    make_pdf(source / "Alpha Manual.pdf", "alpha manual page text")
    return source, tmp_path / "index.sqlite3", tmp_path / "vectors.faiss", StubEmbeddingModel()


def documents_in(database: Path) -> list[str]:
    connection = sqlite3.connect(database)
    try:
        return sorted(row[0] for row in connection.execute("SELECT source_path FROM documents"))
    finally:
        connection.close()


def test_initial_load_indexes_everything_recursively(corpus) -> None:
    source, database, index, model = corpus
    summary = refresh_corpus(source, database, index, model)

    assert summary.discovered == 3
    assert summary.processed == 3
    assert summary.skipped == 0
    assert summary.failed == 0
    assert summary.chunks > 0
    assert summary.embeddings_created == summary.embeddings_total == summary.faiss_vectors
    assert summary.consistent is True
    # A nested folder and both supported formats were picked up.
    assert documents_in(database) == ["Alpha Manual.pdf", "Alpha Process.docx", "nested/Beta Process.docx"]


def test_repeat_load_skips_everything_and_reuses_embeddings(corpus) -> None:
    source, database, index, model = corpus
    first = refresh_corpus(source, database, index, model)
    after_first = model.embedded_texts

    second = refresh_corpus(source, database, index, model)

    assert second.discovered == 3
    assert second.processed == 0
    assert second.skipped == 3
    assert second.embeddings_created == 0
    assert model.embedded_texts == after_first, "unchanged documents must not be re-embedded"
    assert second.embeddings_total == first.embeddings_total
    assert second.faiss_vectors == first.faiss_vectors
    assert second.consistent is True


def test_adding_a_document_processes_only_the_new_one(corpus) -> None:
    source, database, index, model = corpus
    refresh_corpus(source, database, index, model)
    before = model.embedded_texts

    make_docx(source / "Alpha Extra.docx", "Alpha extra", "alpha extra content")
    summary = refresh_corpus(source, database, index, model)

    assert summary.discovered == 4
    assert summary.processed == 1
    assert summary.skipped == 3
    assert summary.embeddings_created == model.embedded_texts - before
    assert summary.embeddings_created > 0
    assert summary.consistent is True
    assert "Alpha Extra.docx" in documents_in(database)


def test_modifying_a_document_replaces_its_chunks_and_embeddings(corpus) -> None:
    source, database, index, model = corpus
    refresh_corpus(source, database, index, model)
    before = model.embedded_texts

    make_docx(source / "Alpha Process.docx", "Alpha heading", "alpha process content, revised and longer")
    summary = refresh_corpus(source, database, index, model)

    assert summary.processed == 1
    assert summary.skipped == 2
    assert summary.embeddings_created == model.embedded_texts - before > 0
    # No orphans: stale chunks and their embeddings went away with the old revision.
    assert summary.consistent is True
    consistency = index_consistency(database, index, model)
    assert consistency["searchable_chunks"] == consistency["embedded_chunks"] == consistency["faiss_vectors"]


def test_removing_a_document_prunes_it_from_the_searchable_index(corpus) -> None:
    source, database, index, model = corpus
    refresh_corpus(source, database, index, model)
    assert any(result["filename"] == "Beta Process.docx"
               for result in search_vector("beta", 10, database, index, model))

    (source / "nested" / "Beta Process.docx").unlink()
    summary = refresh_corpus(source, database, index, model)

    assert summary.pruned == 1
    assert summary.discovered == 2
    assert "nested/Beta Process.docx" not in documents_in(database)
    assert summary.consistent is True
    # The deleted document must no longer be presented to users.
    assert all(result["filename"] != "Beta Process.docx"
               for result in search_vector("beta", 10, database, index, model))


def test_pruning_requires_the_file_to_be_genuinely_absent(tmp_path: Path) -> None:
    """A document missing from the scan but present on disk must never be pruned.

    This is the guard against a transient discovery problem deleting real content.
    """
    source = tmp_path / "corpus"
    make_docx(source / "Present.docx", "Alpha", "alpha")
    indexed = ["Present.docx", "Vanished.docx"]

    prunable = _prunable(source, discovered=set(), indexed=indexed)

    assert prunable == ["Vanished.docx"]
    assert "Present.docx" not in prunable


def test_empty_scan_does_not_wipe_the_index_without_explicit_confirmation(corpus) -> None:
    source, database, index, model = corpus
    refresh_corpus(source, database, index, model)

    for path in list(source.rglob("*.docx")) + list(source.rglob("*.pdf")):
        path.unlink()
    guarded = refresh_corpus(source, database, index, model)

    assert guarded.pruned == 0
    assert guarded.prune_skipped_reason is not None
    assert len(documents_in(database)) == 3

    confirmed = refresh_corpus(source, database, index, model, allow_full_prune=True)
    assert confirmed.pruned == 3
    assert documents_in(database) == []
    # With nothing left to index, the stale vector index is removed rather than left behind.
    assert not index.is_file()
    assert confirmed.consistent is True


def test_large_prune_is_held_back_by_the_safety_limit(corpus) -> None:
    source, database, index, model = corpus
    refresh_corpus(source, database, index, model)

    (source / "Alpha Process.docx").unlink()
    (source / "nested" / "Beta Process.docx").unlink()
    guarded = refresh_corpus(source, database, index, model)

    assert guarded.pruned == 0
    assert "safety limit" in (guarded.prune_skipped_reason or "")
    assert len(documents_in(database)) == 3

    confirmed = refresh_corpus(source, database, index, model, allow_full_prune=True)
    assert confirmed.pruned == 2
    assert confirmed.consistent is True


def test_no_prune_option_retains_removed_documents(corpus) -> None:
    source, database, index, model = corpus
    refresh_corpus(source, database, index, model)

    (source / "nested" / "Beta Process.docx").unlink()
    summary = refresh_corpus(source, database, index, model, prune=False)

    assert summary.pruned == 0
    assert "nested/Beta Process.docx" in documents_in(database)


def test_duplicate_filenames_stay_distinct_documents(tmp_path: Path) -> None:
    """Identity is the source path, not the filename, so neither file is lost."""
    source = tmp_path / "corpus"
    make_docx(source / "one" / "Same Name.docx", "Alpha one", "alpha from folder one")
    make_docx(source / "two" / "Same Name.docx", "Beta two", "beta from folder two")
    database, index, model = tmp_path / "index.sqlite3", tmp_path / "vectors.faiss", StubEmbeddingModel()

    summary = refresh_corpus(source, database, index, model)

    assert summary.processed == 2
    assert documents_in(database) == ["one/Same Name.docx", "two/Same Name.docx"]
    connection = sqlite3.connect(database)
    try:
        identities = [row[0] for row in connection.execute("SELECT document_id FROM documents")]
    finally:
        connection.close()
    assert len(set(identities)) == 2, "same-named files must not collapse to one identity"
    assert summary.consistent is True


def test_duplicate_filenames_are_refused_rather_than_opened_ambiguously(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The open endpoint must never guess which same-named document was meant."""
    from document_finder.api import routes

    source = tmp_path / "corpus"
    make_docx(source / "one" / "Same Name.docx", "Alpha one", "alpha one")
    make_docx(source / "two" / "Same Name.docx", "Beta two", "beta two")
    make_docx(source / "Unique.docx", "Alpha unique", "alpha unique")
    database = tmp_path / "index.sqlite3"
    refresh_corpus(source, database, tmp_path / "vectors.faiss", StubEmbeddingModel())

    monkeypatch.setenv("DOCUMENT_FINDER_DATA_ROOT", str(source))
    monkeypatch.setenv("DOCUMENT_FINDER_DATABASE", str(database))

    assert routes.resolve_document_path("Same Name.docx") is None
    assert routes.resolve_document_path("Unique.docx") == (source / "Unique.docx").resolve()
    # Path traversal stays blocked.
    assert routes.resolve_document_path("../secret.docx") is None
    assert routes.resolve_document_path("one/Same Name.docx") is None


def test_duplicate_content_keeps_both_documents_openable(tmp_path: Path) -> None:
    """Identical bytes under different names must not be deduplicated away."""
    source = tmp_path / "corpus"
    make_docx(source / "Original.docx", "Alpha heading", "alpha identical content")
    (source / "Copy.docx").write_bytes((source / "Original.docx").read_bytes())
    database, index, model = tmp_path / "index.sqlite3", tmp_path / "vectors.faiss", StubEmbeddingModel()

    summary = refresh_corpus(source, database, index, model)

    assert summary.processed == 2
    assert documents_in(database) == ["Copy.docx", "Original.docx"]
    assert summary.consistent is True


def test_failed_ocr_documents_are_only_retried_when_asked(corpus) -> None:
    """A document stored after an OCR failure is otherwise skipped forever."""
    source, database, index, model = corpus
    refresh_corpus(source, database, index, model)

    repository = SQLiteRepository(database)
    try:
        repository.connection.execute(
            "UPDATE documents SET ocr_status = 'failed' WHERE source_path = ?", ("Alpha Manual.pdf",)
        )
        repository.connection.commit()
        assert repository.sources_with_failed_ocr() == ["Alpha Manual.pdf"]
    finally:
        repository.close()

    default_run = refresh_corpus(source, database, index, model)
    assert default_run.processed == 0, "default behaviour is unchanged: the document stays skipped"
    assert default_run.retried_failed_ocr == 0

    retried = refresh_corpus(source, database, index, model, retry_failed_ocr=True)
    assert retried.retried_failed_ocr == 1
    assert retried.processed == 1
    assert retried.consistent is True


def test_consistency_check_detects_a_stale_vector_index(corpus) -> None:
    """Ingesting without rebuilding leaves FAISS behind; the check must say so."""
    source, database, index, model = corpus
    refresh_corpus(source, database, index, model)

    make_docx(source / "Alpha Later.docx", "Alpha later", "alpha later content")
    ingest_only = refresh_corpus(source, database, index, model, build_index=False)

    assert ingest_only.processed == 1
    assert ingest_only.consistent is False
    assert ingest_only.warnings

    repaired = refresh_corpus(source, database, index, model)
    assert repaired.consistent is True


def test_summary_is_serialisable_for_reporting(corpus) -> None:
    source, database, index, model = corpus
    payload = refresh_corpus(source, database, index, model).to_dict()
    for key in ("discovered", "processed", "skipped", "pruned", "failed",
                "chunks", "embeddings_created", "faiss_vectors", "consistent"):
        assert key in payload
