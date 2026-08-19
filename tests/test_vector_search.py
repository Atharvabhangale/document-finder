from __future__ import annotations

import sqlite3
from pathlib import Path

import numpy as np
from docx import Document

from document_finder.ingestion.pipeline import ingest_directory
from document_finder.search.vector import build_vector_index, search_vector


class DummyEmbeddingModel:
    model_name = "test/dummy-embedding"
    configuration = {"version": "1", "normalize": True}

    def __init__(self) -> None:
        self.document_calls = 0

    @staticmethod
    def _vector(text: str) -> list[float]:
        text = text.casefold()
        return [float("procurement" in text), float("change" in text), 0.1]

    def embed_documents(self, texts: list[str]) -> np.ndarray:
        self.document_calls += 1
        return np.asarray([self._vector(text) for text in texts], dtype=np.float32)

    def embed_queries(self, texts: list[str]) -> np.ndarray:
        return np.asarray([self._vector(text) for text in texts], dtype=np.float32)


def _make_doc(path: Path, heading: str, text: str) -> None:
    doc = Document()
    doc.add_heading(heading, level=1)
    doc.add_paragraph(text)
    doc.save(path)


def test_vector_embedding_index_search_aggregation_and_persistence(tmp_path: Path) -> None:
    source = tmp_path / "source"
    (source / "one").mkdir(parents=True)
    (source / "two").mkdir()
    _make_doc(source / "Procurement.docx", "Procurement Kit", "procurement kit creation")
    _make_doc(source / "one" / "Same.docx", "Change A", "change procedure")
    _make_doc(source / "two" / "Same.docx", "Change B", "change procedure")
    database, index = tmp_path / "index.sqlite3", tmp_path / "vectors.faiss"
    assert ingest_directory(source, database).failed == 0
    model = DummyEmbeddingModel()

    first = build_vector_index(database, index, model)
    second = build_vector_index(database, index, model)
    results = search_vector("procurement kit", database_path=database, index_path=index, model=model)
    duplicate_results = search_vector("change", database_path=database, index_path=index, model=model)

    assert first["embedded"] == 3 and first["indexed"] == 3
    assert second["embedded"] == 0 and model.document_calls == 1
    assert index.is_file() and index.with_suffix(".faiss.json").is_file()
    assert results[0]["filename"] == "Procurement.docx"
    assert results[0]["section"] == "Procurement Kit"
    assert results[0]["page"] is None and results[0]["chunk_id"]
    assert [item["filename"] for item in duplicate_results].count("Same.docx") == 1
    connection = sqlite3.connect(database)
    try:
        assert connection.execute("SELECT COUNT(*) FROM chunk_embeddings").fetchone()[0] == 3
    finally:
        connection.close()


def test_vector_empty_or_invalid_query_does_not_load_index(tmp_path: Path) -> None:
    assert search_vector("", database_path=tmp_path / "missing.sqlite3") == []
    assert search_vector(None, database_path=tmp_path / "missing.sqlite3") == []  # type: ignore[arg-type]
