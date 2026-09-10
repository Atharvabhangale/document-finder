"""Phase 18: what the partial-branch experiment actually found.

These tests pin the *discovered* behaviour of the diagnostic and of each
candidate policy, so the reasons the strategy was not adopted are checkable
rather than remembered. They assert the harm as well as the absence of benefit:
if someone later wires a drop-the-dominant-branch policy into production, the
tests here name the queries it breaks and why.

Production behaviour is not exercised here beyond one isolation check —
`tests/test_clarification_decision.py` owns the production decisions, and Phase
18 changed none of them.

Hermetic: the synthetic corpus below is the one from
`tests/test_clarification_decision.py`, reproduced so the two files stay
independent. It mirrors the real 36-document corpus for these queries, and the
real-corpus numbers are recorded in
`docs/phase18-partial-branch-experiment.md`.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pytest

from document_finder.experiments.query_understanding import (
    CorpusProfile,
    DocumentProfile,
    analyze_query,
)
from document_finder.experiments.query_understanding.analyzer import ESCAPE_OPTION
from document_finder.experiments.query_understanding.partial_branch import (
    DUPLICATE,
    SEARCH_ALL,
    STRATEGIES,
    USEFUL,
    WEAK,
    PartialBranchPolicy,
    decide,
    diagnose,
)


CORPUS = CorpusProfile.from_profiles([
    # Change management family (ECR / change management).
    DocumentProfile("BEL_SOP_Change Management Process_PRR ECR ECN_V0.pdf", (), "raise an ecr ecn", "d01"),
    DocumentProfile("PR-ECR-ECN-ACN Process.docx", (), "problem report change request ecr", "d02"),
    DocumentProfile("Change Management Process_.docx", ("Problem Report Proto BOM",),
                    "change management proto bom production bom", "d03"),
    DocumentProfile("User Guide_Company standard document Change Management Process Document.docx",
                    ("Change Request Creation:",), "change notice creation", "d04"),
    DocumentProfile("BEL_SOP_Change Task Deletion & Undo Reservation Activity Process_V0.pdf", (),
                    "delete a change task", "d05"),
    # BOM family.
    DocumentProfile("BEL_SOP_BOM Creation Using Copy Paste Function and Modification of BOM_V0.pdf", (),
                    "bom creation copy paste", "d06"),
    DocumentProfile("BEL_SOP_BOM Creation Using Excel Template Sheet_V0.pdf", (), "bom creation excel", "d07"),
    DocumentProfile("BEL_SOP_Copy BOM Structure from Concept FG Code to Final FG Code_V0.pdf", (),
                    "copy bom structure fg code", "d08"),
    DocumentProfile("Part Creation & EBOM Process.docx", ("Adding BOM Structure", "Significance of 5 Segments"),
                    "part creation ebom bom structure", "d09"),
    # NPD gates (strong enumerable discriminator).
    *[DocumentProfile(f"NPD MS process in PLM - MS{gate} gate (V1).pdf", (), "npd milestone gate", f"g{gate}")
      for gate in range(7)],
    # IT deployment (versioned discriminator, and the BOM cross-domain collision).
    DocumentProfile("Deployment_Document_for_Bajaj_Chatbot_Version_01.docx",
                    ("A) Multi-Level BOM Extraction Deployment",), "java utility deployment", "d10"),
    DocumentProfile("Deployment_Document_for_Bajaj_Chatbot_Version_02.docx", ("Java Utility Deployment",),
                    "java utility deployment", "d11"),
    DocumentProfile("Deployment_Document_for_Bajaj_Chatbot_Version_03.docx", ("Compilation & Server Restart",),
                    "chatbot utility", "d12"),
    # Work request family.
    DocumentProfile("Work Request CRE Combustion Process Document_v2.docx", (), "cre combustion work request", "d13"),
    DocumentProfile("Work Request CRE Process Document - KOEL.docx", (), "cre koel work request", "d14"),
    DocumentProfile("Work Request NVH Process Document.docx", (), "nvh work request", "d15"),
    DocumentProfile("BEL_SOP_NPD Request Creation User Manual_V0.pdf", (), "npd request creation", "d16"),
    DocumentProfile("BEL_SOP_Promotion Request to Released WTPart and CAD Parts_V0.pdf", (),
                    "promotion request wtpart cad", "d17"),
    # Single unambiguous document.
    DocumentProfile("Procurement Kit Process Document.docx", ("Procurement Kit Creation Steps:",),
                    "procurement kit creation", "d18"),
])

PROBES = ("ECR", "work request", "NPD", "BOM", "SOP", "gate", "chatbot", "change management")


def facet(query: str, name: str):
    for candidate in diagnose(query, CORPUS).facets:
        if candidate.facet == name:
            return candidate
    raise AssertionError(f"{query!r} has no {name} facet")


# --------------------------------------------- the diagnostic describes reality


def test_diagnostic_agrees_with_the_analyzers_own_facet_and_options() -> None:
    """The diagnostic rebuilds the facet ladder, so it must not drift from it.

    Whenever the analyzer offers a clarification, the diagnostic's inventory for
    that facet must account for exactly the same options. The one deliberate
    difference is that the diagnostic collapses branches selecting an identical
    document set, which the analyzer does not — so those are added back here.
    """
    for query in PROBES:
        clarification = analyze_query(query, CORPUS).clarification
        if clarification is None:
            continue
        offered = [option for option in clarification.options if option != ESCAPE_OPTION]
        inventory = facet(query, clarification.facet)
        accounted = sorted(
            branch.label for branch in inventory.branches
            if branch.verdict in {WEAK, USEFUL, DUPLICATE}
        )
        if inventory.facet == "distinguishing_term":
            # The analyzer caps this facet at five options, so the inventory is a
            # superset; every offered option must still be accounted for.
            assert set(offered) <= set(accounted), query
        else:
            assert accounted == sorted(offered), query
        assert inventory.candidates_before == clarification.candidates_before, query


def test_every_branch_verdict_is_reproducible_from_its_own_numbers() -> None:
    for query in PROBES:
        report = diagnose(query, CORPUS)
        for inventory in report.facets:
            for branch in inventory.branches:
                if branch.verdict == SEARCH_ALL:
                    assert branch.size == report.candidates_before, (query, branch.label)
                    assert not branch.narrows
                elif branch.verdict in {WEAK, USEFUL}:
                    assert branch.narrows, (query, branch.label)
                    assert (branch.verdict == USEFUL) == (branch.reduction >= 0.5), (query, branch.label)


# ------------------------------------------------------------------------ ECR


def test_ecr_offers_at_most_one_useful_semantic_branch() -> None:
    """ECR is not a threshold problem: the corpus contains no second choice.

    Two documents name ECR. Domain and category cover both (identical to
    searching everything); the document-type facet narrows to one of the two and
    has no second narrowing branch. There is nothing to pick between.
    """
    report = diagnose("ECR", CORPUS)
    assert report.candidates_before == 2
    assert report.analyzer_would_consider is False, "two candidates is below the clarification floor"
    for inventory in report.facets:
        if inventory.facet == "distinguishing_term":
            continue
        assert len(inventory.useful) <= 1, f"{inventory.facet} unexpectedly offers a real choice"


@pytest.mark.parametrize("name", sorted(STRATEGIES))
def test_no_partial_branch_policy_asks_about_ecr(name: str) -> None:
    decision = decide("ECR", CORPUS, STRATEGIES[name])
    assert decision.asks is False
    assert decision.options == ()


def test_forcing_ecr_to_clarify_requires_filename_fragments() -> None:
    """The only facet that could ask about ECR answers with filename fragments.

    The brief's own "bad" example, produced from the corpus rather than asserted:
    the narrowing branches are PR, PRR, ACN, Change and Management. Worse, they
    collapse to just two distinct selections — one per candidate document — so
    the question would be "pick document one or document two", labelled by
    whichever fragment happened to sort first.
    """
    tokens = facet("ECR", "distinguishing_term")
    fragments = {"PR", "PRR", "ACN", "Change", "Management"}
    narrowing = {branch.label for branch in tokens.branches
                 if branch.verdict in {WEAK, USEFUL, DUPLICATE}}
    assert narrowing == fragments
    selections = {branch.documents for branch in tokens.branches if branch.narrows}
    assert len(selections) == 2, "five fragment labels, two actual choices"
    assert all(len(selection) == 1 for selection in selections)


# ---------------------------------------------------------------- work request


def test_work_request_partial_branches_hide_the_on_topic_documents() -> None:
    """Dropping the dominant branch drops the documents the user asked for.

    The category facet's dominant branch *is* "Work Request" — the three
    documents named after the query. Keeping only the useful branches offers the
    two documents that are not work requests and hides the three that are.
    """
    categories = facet("work request", "category")
    assert categories.candidates_before == 5
    dominant = max(categories.narrowing, key=lambda branch: branch.size)
    assert dominant.label == "Work Request" and dominant.size == 3
    assert dominant.verdict == WEAK, "the on-topic branch is what the bar rejects"

    hidden = set(categories.dropped_documents)
    assert hidden == {
        "Work Request CRE Combustion Process Document_v2.docx",
        "Work Request CRE Process Document - KOEL.docx",
        "Work Request NVH Process Document.docx",
    }
    assert categories.useful_coverage == 0.4
    assert [branch.label for branch in categories.useful] == ["NPD / Project", "Part / WTPart / CAD"]


def test_work_requests_useful_only_facets_fail_the_coverage_guard() -> None:
    """Which is why the guarded policy declines rather than asking a bad question."""
    for name in ("category", "document_type"):
        assert facet("work request", name).useful_coverage < 0.5


# ------------------------------------------------------------------------ NPD


def test_npd_partial_branches_are_duplicates_of_each_other() -> None:
    """NPD's two useful type branches select the same single document.

    `BEL_SOP_NPD Request Creation User Manual` is both an SOP and a user manual,
    so "SOP / Work Instruction" and "User Guide / Manual" are one choice shown
    twice — and together they reach 1 of 8 candidates, hiding all seven gates.
    """
    types = facet("NPD", "document_type")
    assert types.candidates_before == 8
    verdicts = {branch.label: branch.verdict for branch in types.branches}
    assert verdicts["Process Document"] == WEAK
    assert DUPLICATE in verdicts.values(), "the second type branch selects the same document"
    assert types.useful_coverage <= 0.125
    assert len(types.dropped_documents) == 7


def test_npds_only_sharply_narrowing_options_are_filename_fragments() -> None:
    """A fragment mixed into an incomplete family: the brief's "bad" pattern.

    Every policy that allows filename tokens answers "NPD" with 'Creation'
    alongside MS0-MS3, leaving MS4, MS5 and MS6 reachable only by searching
    everything. That is worse than the direct search it would replace.
    """
    decision = decide("NPD", CORPUS, STRATEGIES["guarded"])
    assert decision.asks is True
    assert decision.facet == "distinguishing_term"
    assert decision.options == ("Creation", "MS0", "MS1", "MS2", "MS3")
    assert len(decision.dropped_documents) == 3


def test_disallowing_filename_tokens_leaves_npd_on_direct_search() -> None:
    assert decide("NPD", CORPUS, STRATEGIES["semantic_only"]).asks is False


# ------------------------------------------------- regressions the policies cause


@pytest.mark.parametrize("name", ["naive", "guarded", "semantic_only", "phase17_ranked"])
def test_dropping_the_dominant_branch_regresses_bom(name: str) -> None:
    """BOM is the Phase 17 win, and every drop-the-dominant policy undoes it.

    Production asks a full-coverage document-type question. Dropping the
    dominant branch makes the category facet score better on the split while
    offering only the three branches that are *not* about BOM, hiding the three
    BOM SOPs behind "Search everything".

    `phase17_ranked` is the closest possible port of the existing ranking —
    coverage as a bucket rather than a floor — and it regresses BOM too: both
    facets clear the bucket, so the better split wins and the split is better
    precisely because the on-topic branch was dropped.
    """
    production = analyze_query("BOM", CORPUS).clarification
    assert production is not None and production.facet == "document_type"

    decision = decide("BOM", CORPUS, STRATEGIES[name])
    assert decision.asks is True
    assert decision.facet == "category", name
    assert "BOM / EBOM" not in decision.options, name
    assert len(decision.dropped_documents) == 3, name
    assert decision.coverage == 0.5, "exactly at the guard, so the guard does not save it"


def test_unguarded_policy_reintroduces_the_phase_17_coverage_defect() -> None:
    """`SOP`'s document-type facet posts a great split by omitting most candidates.

    Phase 17 chose the full-coverage category facet over it. The naive policy —
    partial branches with no coverage guard — picks it again.
    """
    types = facet("SOP", "document_type")
    assert types.useful_coverage < 0.5 < types.worst_case_reduction
    naive = decide("SOP", CORPUS, PartialBranchPolicy("naive", min_retained_coverage=0.0))
    assert naive.facet == "document_type"
    assert naive.coverage < 0.5
    guarded = decide("SOP", CORPUS, STRATEGIES["guarded"])
    assert guarded.facet == "category" and guarded.coverage == 1.0


@pytest.mark.parametrize("query", ["gate", "chatbot"])
def test_semantic_only_policy_loses_the_best_questions_in_the_corpus(query: str) -> None:
    """MS0-MS6 and Version 01-03 are filename tokens; banning them costs both."""
    assert analyze_query(query, CORPUS).needs_clarification is True
    assert decide(query, CORPUS, STRATEGIES["semantic_only"]).asks is False


def test_change_management_keeps_its_document_type_question_under_every_policy() -> None:
    for name, policy in STRATEGIES.items():
        decision = decide("change management", CORPUS, policy)
        assert decision.asks is True, name
        assert decision.facet == "document_type", name
        assert decision.options == ("SOP / Work Instruction", "User Guide / Manual"), name


# ------------------------------------------------------- invariants of any policy


@pytest.mark.parametrize("name", sorted(STRATEGIES))
def test_no_policy_ever_invents_a_residual_option(name: str) -> None:
    """No "Other", no complement, no escape hatch dressed as a choice.

    Options can only be facet branches, so every label is a corpus-derived
    class name and every option selects at least one real document.
    """
    for query in PROBES:
        decision = decide(query, CORPUS, STRATEGIES[name])
        if not decision.asks:
            continue
        assert ESCAPE_OPTION not in decision.options, (name, query)
        assert not {"Other", "Others", "Everything else", "Miscellaneous"} & set(decision.options)
        assert all(decision.documents), (name, query)
        assert len(decision.options) == len(set(decision.options)), (name, query)


@pytest.mark.parametrize("name", sorted(STRATEGIES))
def test_no_policy_offers_the_same_documents_twice_or_the_whole_set(name: str) -> None:
    for query in PROBES:
        decision = decide(query, CORPUS, STRATEGIES[name])
        if not decision.asks:
            continue
        selections = [frozenset(documents) for documents in decision.documents]
        assert len(set(selections)) == len(selections), (name, query, "duplicate option")
        candidates = diagnose(query, CORPUS).candidates_before
        assert all(0 < len(selection) < candidates for selection in selections), (name, query)


@pytest.mark.parametrize("name", sorted(STRATEGIES))
def test_policies_are_deterministic(name: str) -> None:
    for query in PROBES:
        first = decide(query, CORPUS, STRATEGIES[name])
        assert first == decide(query, CORPUS, STRATEGIES[name]), (name, query)


def test_diagnosis_is_side_effect_free_on_an_unmatched_query() -> None:
    report = diagnose("engine 37 hp 2900 rpm", CORPUS)
    assert report.candidates == () and report.facets == ()
    assert report.analyzer_would_consider is False
    for policy in STRATEGIES.values():
        assert decide("engine 37 hp 2900 rpm", CORPUS, policy).asks is False


# ------------------------------------------------------ the one candidate change


def test_the_only_defensible_new_question_keeps_the_dominant_branch() -> None:
    """`work request` gains a complete, meaningful partition — by *not* dropping.

    This is the one new question the experiment found anywhere on the corpus that
    survives the UX principle: every label is a real document class, the three
    options cover all five candidates, and two of the three narrow by 80%. It is
    the opposite of partial branching — the dominant branch stays on the list.
    """
    decision = decide("work request", CORPUS, STRATEGIES["keep_dominant"])
    assert decision.asks is True
    assert decision.facet == "category"
    assert decision.options == ("NPD / Project", "Part / WTPart / CAD", "Work Request")
    assert decision.coverage == 1.0
    assert decision.dropped_documents == ()
    # And the least helpful answer removes only 40%, which is why production —
    # whose bar is 50% — declines it. No new logic is needed to ask it: the
    # existing QUERY_UNDERSTANDING_MIN_REDUCTION knob covers exactly this case.
    assert decision.worst_case_reduction == 0.4


# ------------------------------------- end to end, on the candidate behaviour


class _StubEmbeddingModel:
    """Bag-of-words vectors over a fixed vocabulary. No model is downloaded."""

    model_name = "test/phase18-stub"
    configuration = {"version": "1", "normalize": True}
    VOCABULARY = ("work", "request", "npd", "promotion", "cre", "nvh", "koel")

    @classmethod
    def _vector(cls, text: str) -> list[float]:
        lowered = (text or "").casefold()
        # A small constant keeps every chunk retrievable, so scoping a result set
        # is what narrows it — never an empty index.
        return [1.0 if word in lowered else 0.0 for word in cls.VOCABULARY] + [0.1]

    def embed_documents(self, texts: list[str]) -> "np.ndarray":
        return np.asarray([self._vector(text) for text in texts], dtype=np.float32)

    def embed_queries(self, texts: list[str]) -> "np.ndarray":
        return np.asarray([self._vector(text) for text in texts], dtype=np.float32)


@pytest.fixture
def dominant_branch_corpus(monkeypatch: pytest.MonkeyPatch, tmp_path: Path):
    """A real indexed corpus shaped like `work request`.

    Three documents named after the query (the dominant branch) plus two that
    merely share the word "request" and fall in other categories — the exact
    structure the experiment found on the real corpus.
    """
    pytest.importorskip("docx")
    pytest.importorskip("faiss")
    from docx import Document
    from fastapi.testclient import TestClient

    from document_finder import config, understanding
    from document_finder.api import routes
    from document_finder.api.app import app
    from document_finder.corpus import refresh_corpus
    from document_finder.search.vector import search_vector

    def write(relative: str, heading: str, body: str) -> Path:
        path = tmp_path / "corpus" / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        document = Document()
        document.add_heading(heading, level=1)
        document.add_paragraph(body)
        document.save(path)
        return path

    write("Work Request CRE Combustion Process Document.docx",
          "Work Request CRE Combustion", "cre combustion work request")
    write("Work Request CRE Process Document - KOEL.docx",
          "Work Request CRE KOEL", "cre koel work request")
    write("Work Request NVH Process Document.docx",
          "Work Request NVH", "nvh work request")
    write("BEL_SOP_NPD Request Creation User Manual.docx",
          "NPD Request Creation", "npd request creation")
    write("BEL_SOP_Promotion Request to Released WTPart and CAD Parts.docx",
          "Promotion Request", "promotion request wtpart cad")
    # Two same-named documents in different folders, to prove identity holds.
    write("Engineering/procedure.docx", "Engineering procedure", "work request engineering")
    write("Quality/procedure.docx", "Quality procedure", "work request quality")

    root = tmp_path / "corpus"
    monkeypatch.setenv(config.DATA_ROOT_VARIABLE, str(root))
    monkeypatch.delenv(config.DATABASE_VARIABLE, raising=False)
    monkeypatch.delenv(config.INDEX_VARIABLE, raising=False)
    monkeypatch.delenv(config.QUERY_UNDERSTANDING_VARIABLE, raising=False)
    # The candidate behaviour, reached through the existing configuration knob:
    # `work request`'s complete partition has a 40% worst case.
    monkeypatch.setenv(config.MIN_REDUCTION_VARIABLE, "0.4")
    understanding.clear_cache()

    model = _StubEmbeddingModel()
    refresh_corpus(root, model=model)
    monkeypatch.setattr(
        routes, "search_vector",
        lambda query, **kwargs: search_vector(
            query, kwargs.get("limit", 5), config.database_path(), config.index_path(), model
        ),
    )
    return TestClient(app), root


def test_candidate_question_runs_the_whole_chain(dominant_branch_corpus) -> None:
    """query -> understanding -> option -> scoped search -> document id -> file.

    Verifies the mechanism the candidate behaviour would need, so the decision
    not to adopt it rests on the UX evidence and not on unproven plumbing.
    """
    client, root = dominant_branch_corpus

    understood = client.post("/query-understanding", json={"query": "work request"}).json()
    assert understood["needs_clarification"] is True
    assert understood["question"]
    labels = [option["label"] for option in understood["options"]]
    # The complete corpus-derived partition, dominant branch included.
    assert labels == ["NPD / Project", "Part / WTPart / CAD", "Work Request"]

    unrestricted = client.post("/search", json={"query": "work request", "limit": 10}).json()
    assert unrestricted["results"]

    from document_finder import understanding as layer

    for label in labels:
        narrowed = client.post(
            "/search", json={"query": "work request", "limit": 10, "clarification": label}
        ).json()["results"]
        assert narrowed, f"{label} must not empty the result set"
        assert len(narrowed) < len(unrestricted["results"]), f"{label} must narrow"
        chosen = set(layer.documents_for_choice("work request", label) or ())
        assert {result["document_id"] for result in narrowed} <= chosen

        # Every narrowed result opens its own original file, byte for byte.
        for result in narrowed:
            response = client.get(f"/documents/by-id/{result['document_id']}")
            assert response.status_code == 200, result["filename"]
            folder = root / result["folder"] if result["folder"] else root
            assert response.content == (folder / result["filename"]).read_bytes()


def test_candidate_question_preserves_search_all_and_query_change(dominant_branch_corpus) -> None:
    client, _ = dominant_branch_corpus
    plain = client.post("/search", json={"query": "work request", "limit": 10}).json()
    escaped = client.post(
        "/search", json={"query": "work request", "limit": 10, "clarification": None}
    ).json()
    assert plain == escaped, "Search all documents must equal searching without the layer"

    scoped = client.post(
        "/search", json={"query": "work request", "limit": 10, "clarification": "Work Request"}
    ).json()
    assert scoped != plain, "an answer must actually scope the results"

    # Changing the query re-decides from scratch: an answer to the old question
    # means nothing to the new one and must not filter it.
    other = {"query": "promotion request", "limit": 10}
    assert (client.post("/search", json={**other, "clarification": "Work Request"}).json()
            == client.post("/search", json=other).json())


def test_duplicate_filenames_stay_distinct_under_the_candidate_question(
    dominant_branch_corpus,
) -> None:
    client, root = dominant_branch_corpus
    results = client.post("/search", json={"query": "work request", "limit": 20}).json()["results"]
    duplicates = [result for result in results if result["filename"] == "procedure.docx"]
    assert len(duplicates) == 2, "two same-named documents must stay two results"
    assert {result["folder"] for result in duplicates} == {"Engineering", "Quality"}
    assert len({result["document_id"] for result in duplicates}) == 2
    for result in duplicates:
        response = client.get(f"/documents/by-id/{result['document_id']}")
        assert response.status_code == 200
        assert response.content == (root / result["folder"] / result["filename"]).read_bytes()


def test_disabling_the_layer_still_suppresses_the_candidate_question(
    dominant_branch_corpus, monkeypatch: pytest.MonkeyPatch
) -> None:
    from document_finder import config

    client, _ = dominant_branch_corpus
    monkeypatch.setenv(config.QUERY_UNDERSTANDING_VARIABLE, "false")
    assert client.post("/query-understanding", json={"query": "work request"}).json() == {
        "query": "work request", "needs_clarification": False, "question": None, "options": [],
    }


def test_default_configuration_leaves_work_request_on_direct_search(
    dominant_branch_corpus, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The shipped default is unchanged: only the explicit 0.4 override asks."""
    from document_finder import config, understanding as layer

    client, _ = dominant_branch_corpus
    monkeypatch.delenv(config.MIN_REDUCTION_VARIABLE, raising=False)
    layer.clear_cache()
    assert config.min_clarification_reduction() == 0.5
    assert client.post(
        "/query-understanding", json={"query": "work request"}
    ).json()["needs_clarification"] is False


# ------------------------------------------------ latent finding, not a live bug


def test_analyzer_can_build_duplicate_options_but_never_asks_them() -> None:
    """Found while diagnosing NPD; characterised here rather than fixed.

    `analyzer._finalize` keeps every branch that selects a strict, non-empty
    subset, without noticing that two branches may select the *same* subset.
    `NPD` is such a case: its document-type facet offers "SOP / Work Instruction"
    and "User Guide / Manual" for one and the same document.

    No user can see this today — every clarification containing a duplicate is
    declined by the reduction bar before it is asked — so it is recorded, not
    repaired: fixing it would change production for no observable gain. This
    test fails if that ever stops being true.
    """
    for query in PROBES:
        clarification = analyze_query(query, CORPUS).clarification
        if clarification is None:
            continue
        selections = [
            frozenset(documents) for option, documents in clarification.option_documents.items()
            if option != ESCAPE_OPTION
        ]
        duplicated = len(selections) - len(set(selections))
        if not duplicated:
            continue
        assert query == "NPD", f"a new query now builds duplicate options: {query!r}"
        reduction = 1 - clarification.largest_option_candidates / clarification.candidates_before
        assert reduction < 0.5, "a question with duplicate options must not reach the user"


# ------------------------------------------------------------------- isolation


def test_prototype_is_not_wired_into_production() -> None:
    """Phase 18 must stay an experiment, exactly as Phase 13 did."""
    from document_finder.experiments.query_understanding import partial_branch

    source = Path(partial_branch.__file__).read_text(encoding="utf-8")
    assert "document_finder.search" not in source
    assert "faiss" not in source
    assert "document_finder.understanding" not in source, (
        "the prototype must not reach into the production decision layer"
    )

    production = Path("src/document_finder")
    for relative in ("understanding.py", "config.py", "api/routes.py", "api/app.py",
                     "search/vector.py", "search/lexical.py", "search/hybrid.py",
                     "ingestion/pipeline.py"):
        text = (production / relative).read_text(encoding="utf-8")
        assert "partial_branch" not in text, f"{relative} must not import the Phase 18 prototype"

    exported = Path("src/document_finder/experiments/query_understanding/__init__.py").read_text(
        encoding="utf-8"
    )
    assert "partial_branch" not in exported, (
        "the prototype is imported explicitly, never re-exported alongside the analyzer"
    )
