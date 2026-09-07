"""Query-understanding integration: the optional layer in front of search (Phase 16).

Hermetic throughout — synthetic corpora in temporary folders and a stub
embedding model, so nothing here reads the 36-document corpus or downloads a
model. The Phase 13 analyzer and its frozen dataset are used, never modified.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pytest
from docx import Document
from fastapi.testclient import TestClient

from document_finder import config, understanding
from document_finder.api import routes
from document_finder.api.app import app
from document_finder.corpus import refresh_corpus
from document_finder.search.vector import search_vector


client = TestClient(app)


class StubEmbeddingModel:
    model_name = "test/stub-embedding"
    configuration = {"version": "1", "normalize": True}

    @staticmethod
    def _vector(text: str) -> list[float]:
        text = text.casefold()
        return [
            float("gate" in text or "milestone" in text),
            float("sop" in text or "procedure" in text),
            float("ecr" in text or "change" in text or "procurement" in text) or 0.1,
        ]

    def embed_documents(self, texts: list[str]) -> np.ndarray:
        return np.asarray([self._vector(text) for text in texts], dtype=np.float32)

    def embed_queries(self, texts: list[str]) -> np.ndarray:
        return np.asarray([self._vector(text) for text in texts], dtype=np.float32)


def make_docx(path: Path, heading: str, body: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(body)
    document.save(path)


@pytest.fixture
def corpus(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Path:
    """A corpus holding both a strong and a deliberately weak discriminator."""
    root = tmp_path / "corpus"
    # Strong discriminator: an enumerable token unique to each document.
    for gate in range(6):
        make_docx(root / f"NPD MS process in PLM - MS{gate} gate.docx",
                  f"MS{gate} milestone gate", f"npd milestone gate ms{gate}")
    # Weak discriminator: a skewed split — five PLM SOPs against one IT SOP, so
    # the only available facet barely narrows anything.
    for name in ("Folder Navigation", "Project Creation", "Promotion Request",
                 "Change Task Deletion", "Imported Parts"):
        make_docx(root / f"BEL_SOP_{name}.docx", name, f"{name.casefold()} sop procedure")
    make_docx(root / "Iraje_PAM_SOP.docx", "Iraje PAM", "iraje pam privileged access sop")
    # Unambiguous single documents.
    make_docx(root / "Procurement Kit Process Document.docx", "Procurement Kit", "procurement kit creation")
    make_docx(root / "BEL_SOP_Change Management Process_PRR ECR ECN.docx",
              "Change Management", "raise an ecr and ecn change request")

    monkeypatch.setenv(config.DATA_ROOT_VARIABLE, str(root))
    monkeypatch.delenv(config.DATABASE_VARIABLE, raising=False)
    monkeypatch.delenv(config.INDEX_VARIABLE, raising=False)
    monkeypatch.delenv(config.QUERY_UNDERSTANDING_VARIABLE, raising=False)
    monkeypatch.delenv(config.MIN_REDUCTION_VARIABLE, raising=False)
    understanding.clear_cache()

    model = StubEmbeddingModel()
    refresh_corpus(root, model=model)
    monkeypatch.setattr(
        routes, "search_vector",
        lambda query, **kwargs: search_vector(
            query, kwargs.get("limit", 5), config.database_path(), config.index_path(), model
        ),
    )
    return root


# ------------------------------------------------------------- decision layer


def test_clearly_procedural_query_goes_straight_to_search(corpus: Path) -> None:
    prepared = understanding.prepare("how do I create an ECR?")
    assert prepared.needs_clarification is False
    assert prepared.reason.startswith("direct")


def test_clearly_specific_query_goes_straight_to_search(corpus: Path) -> None:
    prepared = understanding.prepare("procurement")
    assert prepared.needs_clarification is False


def test_strong_discriminator_asks_a_question(corpus: Path) -> None:
    prepared = understanding.prepare("gate")
    assert prepared.needs_clarification is True
    assert prepared.question
    assert len(prepared.options) >= 2
    # Every offered option must select a strict, non-empty subset.
    assert all(option.document_ids for option in prepared.options)
    assert prepared.reduction >= config.min_clarification_reduction()


def test_weak_discriminator_does_not_ask_a_pointless_question(monkeypatch: pytest.MonkeyPatch) -> None:
    """A facet that barely narrows anything must not cost the user a round trip.

    The corpus here offers exactly one viable facet and it is deliberately weak:
    six documents sharing a domain and a document type, split 5/1 by category.
    That is a real choice the analyzer will offer, and a 17% worst-case reduction
    the integration layer must still decline.
    """
    from document_finder.experiments.query_understanding import CorpusProfile, DocumentProfile

    monkeypatch.delenv(config.QUERY_UNDERSTANDING_VARIABLE, raising=False)
    monkeypatch.delenv(config.MIN_REDUCTION_VARIABLE, raising=False)
    weak = CorpusProfile.from_profiles([
        *[DocumentProfile(f"Widget Part Handling {index}.docx", (), "widget part handling", f"w{index}")
          for index in range(5)],
        DocumentProfile("Widget Procurement Kit.docx", (), "widget procurement kit", "w9"),
    ])
    monkeypatch.setattr(understanding, "corpus_profile", lambda *a, **k: weak)

    prepared = understanding.prepare("widget")

    assert prepared.facet == "category", "the weak category facet is the only one available"
    assert prepared.candidates_before == 6 and prepared.largest_option_candidates == 5
    assert prepared.reduction < config.min_clarification_reduction()
    assert prepared.needs_clarification is False
    assert prepared.reason == "reduction_too_small"


def test_out_of_corpus_query_does_not_ask_a_question(corpus: Path) -> None:
    prepared = understanding.prepare("how do I file a travel expense claim?")
    assert prepared.needs_clarification is False
    assert prepared.options == ()


@pytest.mark.parametrize("query", ["", "   ", "???", None])
def test_invalid_queries_never_ask_a_question(corpus: Path, query: object) -> None:
    prepared = understanding.prepare(query)  # type: ignore[arg-type]
    assert prepared.needs_clarification is False


def test_threshold_is_configurable(corpus: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Lowering the bar lets the weak discriminator through; raising it blocks the strong one."""
    monkeypatch.setenv(config.MIN_REDUCTION_VARIABLE, "0.0")
    assert understanding.prepare("sop").needs_clarification is True
    monkeypatch.setenv(config.MIN_REDUCTION_VARIABLE, "0.99")
    assert understanding.prepare("gate").needs_clarification is False


def test_feature_flag_disables_the_layer_entirely(corpus: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv(config.QUERY_UNDERSTANDING_VARIABLE, "false")
    prepared = understanding.prepare("gate")
    assert prepared.needs_clarification is False
    assert prepared.reason == "disabled"
    assert client.post("/query-understanding", json={"query": "gate"}).json() == {
        "query": "gate", "needs_clarification": False, "question": None, "options": [],
    }


def test_analyzer_failure_falls_back_to_search(corpus: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """The experimental layer must never be able to break search."""
    def explode(*args: object, **kwargs: object) -> None:
        raise RuntimeError("analyzer exploded")

    monkeypatch.setattr(understanding, "analyze_query", explode)
    prepared = understanding.prepare("gate")
    assert prepared.needs_clarification is False
    assert prepared.reason == "unavailable"

    response = client.post("/query-understanding", json={"query": "gate"})
    assert response.status_code == 200
    assert response.json()["needs_clarification"] is False
    # And a normal search still works.
    assert client.post("/search", json={"query": "gate"}).status_code == 200


def test_missing_index_does_not_break_the_layer(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    monkeypatch.setenv(config.DATA_ROOT_VARIABLE, str(tmp_path / "nothing-here"))
    monkeypatch.delenv(config.DATABASE_VARIABLE, raising=False)
    understanding.clear_cache()
    prepared = understanding.prepare("gate")
    assert prepared.needs_clarification is False
    assert prepared.reason == "unavailable"


# ----------------------------------------------------------------------- API


def test_clarification_response_shape(corpus: Path) -> None:
    payload = client.post("/query-understanding", json={"query": "gate"}).json()
    assert payload["query"] == "gate"
    assert payload["needs_clarification"] is True
    assert payload["question"]
    assert len(payload["options"]) >= 2
    assert all({"label", "value"} == set(option) for option in payload["options"])
    # No internal concepts are exposed to the client.
    assert not {"intent", "confidence", "ambiguity", "candidate_count", "candidates_before",
                "document_ids", "facet"} & set(payload)


def test_direct_search_response_shape(corpus: Path) -> None:
    payload = client.post("/query-understanding", json={"query": "how do I create an ECR?"}).json()
    assert payload["needs_clarification"] is False
    assert payload["question"] is None and payload["options"] == []


def test_empty_query_is_answered_without_clarification(corpus: Path) -> None:
    payload = client.post("/query-understanding", json={"query": "   "}).json()
    assert payload["needs_clarification"] is False


# --------------------------------------------------------------- integration


def test_selecting_an_option_narrows_the_results(corpus: Path) -> None:
    prepared = client.post("/query-understanding", json={"query": "gate"}).json()
    unrestricted = client.post("/search", json={"query": "gate", "limit": 10}).json()["results"]

    option = prepared["options"][0]
    narrowed = client.post(
        "/search", json={"query": "gate", "limit": 10, "clarification": option["value"]}
    ).json()["results"]

    assert len(narrowed) < len(unrestricted), "answering the question must narrow the result set"
    assert narrowed, "narrowing must not empty the results"
    # Everything shown is genuinely part of the chosen option.
    chosen = set(understanding.documents_for_choice("gate", option["value"]) or ())
    assert {result["document_id"] for result in narrowed} <= chosen
    # Result shape is preserved.
    for result in narrowed:
        assert {"document_id", "filename", "folder", "score", "section", "page"} == set(result)


def test_search_all_reproduces_normal_search(corpus: Path) -> None:
    """The escape route must be byte-identical to searching without the layer."""
    plain = client.post("/search", json={"query": "gate", "limit": 10}).json()
    escaped = client.post("/search", json={"query": "gate", "limit": 10, "clarification": None}).json()
    assert plain == escaped


def test_unknown_or_stale_choice_does_not_narrow_or_fail(corpus: Path) -> None:
    plain = client.post("/search", json={"query": "gate", "limit": 10}).json()
    stale = client.post(
        "/search", json={"query": "gate", "limit": 10, "clarification": "No Such Option"}
    ).json()
    assert stale == plain


def test_choice_on_a_direct_query_is_ignored(corpus: Path) -> None:
    plain = client.post("/search", json={"query": "procurement", "limit": 10}).json()
    with_choice = client.post(
        "/search", json={"query": "procurement", "limit": 10, "clarification": "Anything"}
    ).json()
    assert with_choice == plain


def test_changing_the_query_starts_a_fresh_decision(corpus: Path) -> None:
    assert client.post("/query-understanding", json={"query": "gate"}).json()["needs_clarification"] is True
    assert client.post(
        "/query-understanding", json={"query": "how do I create an ECR?"}
    ).json()["needs_clarification"] is False


# ----------------------------------------------------------------- regression


def test_normal_search_is_unchanged_when_the_layer_is_disabled(
    corpus: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    enabled = client.post("/search", json={"query": "gate", "limit": 10}).json()
    monkeypatch.setenv(config.QUERY_UNDERSTANDING_VARIABLE, "false")
    disabled = client.post("/search", json={"query": "gate", "limit": 10}).json()
    assert enabled == disabled, "/search must not depend on the clarification layer"


def test_document_opening_and_duplicate_filenames_still_work(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    root = tmp_path / "corpus"
    make_docx(root / "Engineering" / "procedure.docx", "Alpha engineering", "gate procedure engineering")
    make_docx(root / "Quality" / "procedure.docx", "Alpha quality", "gate procedure quality")
    monkeypatch.setenv(config.DATA_ROOT_VARIABLE, str(root))
    monkeypatch.delenv(config.DATABASE_VARIABLE, raising=False)
    understanding.clear_cache()
    model = StubEmbeddingModel()
    refresh_corpus(root, model=model)
    monkeypatch.setattr(
        routes, "search_vector",
        lambda query, **kwargs: search_vector(
            query, kwargs.get("limit", 5), config.database_path(), config.index_path(), model
        ),
    )

    results = client.post("/search", json={"query": "procedure", "limit": 10}).json()["results"]
    assert len({result["document_id"] for result in results}) == 2
    assert {result["folder"] for result in results} == {"Engineering", "Quality"}
    for result in results:
        opened = client.get(f"/documents/by-id/{result['document_id']}")
        assert opened.status_code == 200
        assert opened.content == (root / result["folder"] / result["filename"]).read_bytes()


def test_corpus_profile_is_cached_and_refreshes_on_reindex(corpus: Path) -> None:
    understanding.clear_cache()
    first = understanding.corpus_profile()
    assert understanding.corpus_profile() is first, "profile must be cached per request"

    make_docx(corpus / "Extra Gate.docx", "MS9 milestone gate", "npd milestone gate ms9")
    refresh_corpus(corpus, model=StubEmbeddingModel())
    reloaded = understanding.corpus_profile()
    assert reloaded is not first, "a re-index must invalidate the cache"
    assert len(reloaded.documents) == len(first.documents) + 1
